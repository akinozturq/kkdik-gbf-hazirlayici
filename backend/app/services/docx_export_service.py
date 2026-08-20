import io
import os
import zipfile
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.services.translation_service import translation_service
from app.services.reference_service import reference_service

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "GBF_ANTET_TEMPLATE.docx")
FONT_NAME = "Google Sans"

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, color="444444", sz="4", val="single"):
    """Sets clean border around a cell without background fill."""
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

class DocxExportService:
    """
    Generates professional 16-section KKDİK / REACH Annex II compliant Word (.docx) documents
    with strictly separated subheadings according to REACH Annex II / KKDİK Ek-2,
    Google Sans typography, and exact font sizes (12pt main, 11pt sub, 10pt normal text & table headers, 8pt table contents).
    """

    @classmethod
    def generate_docx(cls, product_dict: dict, lang: str = "tr") -> io.BytesIO:
        sds = product_dict.get("sds_data") or {}
        meta = sds.get("meta") or {}
        lang_clean = (lang or "tr").lower()
        t = translation_service.get_sections(lang_clean)

        urun_adi = product_dict.get("urun_adi") or "KİMYASAL ÜRÜN"
        ticari_kod = product_dict.get("ticari_kod") or "KOD-001"
        hazirlama_tarihi = meta.get("hazirlama_tarihi") or "01.01.2026"
        revizyon_no = str(meta.get("revizyon_no") or "00")
        revizyon_tarihi = meta.get("revizyon_tarihi") or hazirlama_tarihi

        with open(TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()

        zip_in = zipfile.ZipFile(io.BytesIO(template_bytes), "r")
        zip_buffer = io.BytesIO()
        zip_out = zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED)

        for item in zip_in.infolist():
            content = zip_in.read(item.filename)
            if item.filename in ("word/header1.xml", "word/header2.xml"):
                text = content.decode("utf-8")
                text = re.sub(r'<w:t>16\.03\.2019</w:t>', f'<w:t>{hazirlama_tarihi}</w:t>', text)
                text = re.sub(r'<w:t>01</w:t>', f'<w:t>{revizyon_no}</w:t>', text)
                text = re.sub(r'<w:t>24\.06\.2026</w:t>', f'<w:t>{revizyon_tarihi}</w:t>', text)
                if lang_clean == "en":
                    text = text.replace("GÜVENLİK BİLGİ FORMU", "SAFETY DATA SHEET")
                    text = text.replace("Hazırlama Tarihi", "Compilation Date")
                    text = text.replace("Yenileme Tarihi", "Revision Date")
                    text = text.replace("Revizyon No", "Revision No")
                content = text.encode("utf-8")
            zip_out.writestr(item, content)

        zip_in.close()
        zip_out.close()
        zip_buffer.seek(0)

        doc = Document(zip_buffer)

        p_elements = list(doc.element.body)
        for p in p_elements:
            if p.tag.endswith("p") or p.tag.endswith("tbl"):
                if not p.tag.endswith("sectPr"):
                    p.getparent().remove(p)

        cls._add_section_1(doc, sds, product_dict, t, lang_clean)
        cls._add_section_2(doc, sds, t, lang_clean)
        cls._add_section_3(doc, sds, t, lang_clean)
        cls._add_section_4(doc, sds, t, lang_clean)
        cls._add_section_5(doc, sds, t, lang_clean)
        cls._add_section_6(doc, sds, t, lang_clean)
        cls._add_section_7(doc, sds, t, lang_clean)
        cls._add_section_8(doc, sds, t, lang_clean)
        cls._add_section_9(doc, sds, t, lang_clean)
        cls._add_section_10(doc, sds, t, lang_clean)
        cls._add_section_11(doc, sds, t, lang_clean)
        cls._add_section_12(doc, sds, t, lang_clean)
        cls._add_section_13(doc, sds, t, lang_clean)
        cls._add_section_14(doc, sds, t, lang_clean)
        cls._add_section_15(doc, sds, t, lang_clean)
        cls._add_section_16(doc, sds, t, lang_clean)

        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        return output_stream

    @classmethod
    def _add_section_banner(cls, doc, title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.font.name = FONT_NAME
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        pPr = p._element.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="8" w:space="2" w:color="000000"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

    @classmethod
    def _add_subsection_title(cls, doc, title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.name = FONT_NAME
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    @classmethod
    def _add_kv_line(cls, doc, label: str, val: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15

        run_lbl = p.add_run(f"{label}: ")
        run_lbl.font.name = FONT_NAME
        run_lbl.font.size = Pt(10)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = RGBColor(0, 0, 0)

        run_val = p.add_run(str(val or "—"))
        run_val.font.name = FONT_NAME
        run_val.font.size = Pt(10)
        run_val.font.color.rgb = RGBColor(20, 20, 20)

    @classmethod
    def _add_text_line(cls, doc, text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(str(text or "—"))
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(20, 20, 20)

    @classmethod
    def _add_section_1(cls, doc, sds, prod, t, lang):
        cls._add_section_banner(doc, t.get("s1", {}).get("title") or "1. MADDENİN / KARIŞIMIN VE ŞİRKETİN / DAĞITICININ KİMLİĞİ")
        b1 = sds.get("b1_kimlik") or {}
        b1_1 = b1.get("b1_1") or {}
        b1_2 = b1.get("b1_2") or {}
        b1_3 = b1.get("b1_3") or {}
        b1_4 = b1.get("b1_4") or {}

        s1_t = t.get("s1", {})
        cls._add_subsection_title(doc, s1_t.get("s1_1") or "1.1. Madde / Karışım Kimliği")
        cls._add_kv_line(doc, s1_t.get("product_name") or "Madde / Karışım Adı", b1_1.get("madde_karisim_adi") or prod.get("urun_adi"))
        cls._add_kv_line(doc, s1_t.get("trade_code") or "Ticari Kod / Stok Kodu", prod.get("ticari_kod"))
        cls._add_kv_line(doc, s1_t.get("reach_reg") or "Kayıt Numarası", b1_1.get("kayit_numarasi") or ("Exempt from registration / Not applicable." if lang == "en" else "Kayıttan muaftır / Uygulanabilir değildir."))

        cls._add_subsection_title(doc, s1_t.get("s1_2") or "1.2. Madde veya Karışımın Belirlenmiş Kullanımları ve Tavsiye Edilmeyen Kullanımları")
        cls._add_kv_line(doc, s1_t.get("identified_uses") or "Belirlenmiş Kullanımlar", ", ".join(b1_2.get("tanimlanmis_kullanimlar") or []) or ("Industrial / Professional use" if lang == "en" else "Sanayi / Endüstriyel kullanım"))
        cls._add_kv_line(doc, s1_t.get("uses_advised_against") or "Tavsiye Edilmeyen Kullanımlar", ", ".join(b1_2.get("tavsiye_edilmeyen_kullanimlar") or []) or ("Do not use for purposes other than those identified." if lang == "en" else "Tavsiye edilen kullanımların dışında kullanılmamalıdır."))

        cls._add_subsection_title(doc, s1_t.get("s1_3") or "1.3. Güvenlik Bilgi Formu Tedarikçisinin Bilgileri")
        cls._add_kv_line(doc, s1_t.get("supplier_name") or "Tedarikçi / İmalatçı Adı", b1_3.get("tedarikci_adi") or "Aypol Kimya Sanayi ve Ticaret A.Ş.")
        cls._add_kv_line(doc, s1_t.get("address") or "Adres", b1_3.get("adres") or ("Istanbul / Turkey" if lang == "en" else "İstanbul / Türkiye"))
        cls._add_kv_line(doc, s1_t.get("phone") or "İletişim", f"Tel: {b1_3.get('telefon') or '—'} | E-Mail: {b1_3.get('eposta') or '—'}")
        cls._add_kv_line(doc, s1_t.get("email") or "Yetkili Kişi (KDU)", b1_3.get("yetkili_kisi") or "—")

        cls._add_subsection_title(doc, s1_t.get("s1_4") or "1.4. Acil Durum Telefon Numarası")
        cls._add_kv_line(doc, s1_t.get("emergency_phone") or "Acil Durum Telefonu", b1_4.get("acil_telefon") or ("112 Emergency / National Poison Centre" if lang == "en" else "114 (UZEM - Ulusal Zehir Danışma Merkezi) | 112 Acil"))

    @classmethod
    def _add_section_2(cls, doc, sds, t, lang):
        s2_t = t.get("s2", {})
        cls._add_section_banner(doc, s2_t.get("title") or "2. ZARARLILIK TANIMI")
        b2 = sds.get("b2_zarar_tanimi") or {}
        b2_1 = b2.get("b2_1") or {}
        b2_2 = b2.get("b2_2") or {}
        b2_3 = b2.get("b2_3") or {}

        cls._add_subsection_title(doc, s2_t.get("s2_1") or "2.1. Maddenin veya Karışımın Sınıflandırılması")
        siniflar = b2_1.get("siniflandirmalar") or []
        if siniflar:
            for s in siniflar:
                cls._add_text_line(doc, f"• {s.get('zararlilik_sinifi', '')} ({s.get('kategori', '')}): {s.get('h_kodu', '')}")
        else:
            cls._add_text_line(doc, b2_1.get("siniflandirilmama_gerekcesi") or (s2_t.get("not_classified") or "SEA Yönetmeliği kriterlerine göre zararlı olarak sınıflandırılmamıştır."))

        cls._add_subsection_title(doc, s2_t.get("s2_2") or "2.2. Etiket Unsurları")
        
        piks = b2_2.get("piktogramlar") or []
        p_pik = doc.add_paragraph()
        p_pik.paragraph_format.space_before = Pt(2)
        p_pik.paragraph_format.space_after = Pt(3)
        p_pik.paragraph_format.line_spacing = 1.15

        run_lbl = p_pik.add_run(f"{s2_t.get('pictograms') or 'Piktogramlar'}: ")
        run_lbl.font.name = FONT_NAME
        run_lbl.font.size = Pt(10)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = RGBColor(0, 0, 0)

        if piks:
            pictograms_dir = os.path.join(os.path.dirname(__file__), "..", "templates", "pictograms")
            for pic_code in piks:
                if not pic_code:
                    continue
                p_code = pic_code.strip()
                png_candidates = [
                    os.path.join(pictograms_dir, f"{p_code.upper()}.png"),
                    os.path.join(pictograms_dir, f"{p_code.lower()}.png"),
                    os.path.join("pictograms", "png", f"{p_code.upper()}.png"),
                    os.path.join("pictograms", "png", f"{p_code.lower()}.png"),
                ]
                found_path = next((p for p in png_candidates if os.path.exists(p)), None)
                if found_path:
                    run_img = p_pik.add_run()
                    run_img.add_picture(found_path, width=Inches(0.65), height=Inches(0.65))
                    p_pik.add_run("  ")
                else:
                    run_txt = p_pik.add_run(f"[{p_code}] ")
                    run_txt.font.name = FONT_NAME
                    run_txt.font.size = Pt(10)
        else:
            run_none = p_pik.add_run(s2_t.get("no_pictogram") or "GHS Piktogramı Bulunmamaktadır")
            run_none.font.name = FONT_NAME
            run_none.font.size = Pt(10)

        raw_sig = b2_2.get("uyari_kelimesi") or "Yok"
        translated_sig = translation_service.translate_signal_word(raw_sig, lang)
        cls._add_kv_line(doc, s2_t.get("signal_word") or "Uyarı Kelimesi", translated_sig)
        cls._add_kv_line(doc, s2_t.get("hazard_statements") or "Zararlılık İfadeleri (H)", ", ".join(b2_2.get("h_ifadeleri") or []) or "—")
        cls._add_kv_line(doc, s2_t.get("precautionary_statements") or "Önlem İfadeleri (P)", ", ".join(b2_2.get("p_ifadeleri") or []) or "—")

        cls._add_subsection_title(doc, s2_t.get("s2_3") or "2.3. Diğer Zararlar")
        cls._add_text_line(doc, b2_3.get("pbt_vpvb_degerlendirme") or (s2_t.get("pbt_vpvb") or "PBT / vPvB kriterlerini karşılamaz."))

    @classmethod
    def _add_section_3(cls, doc, sds, t, lang):
        s3_t = t.get("s3", {})
        cls._add_section_banner(doc, s3_t.get("title") or "3. BİLEŞİMİ / İÇİNDEKİLER HAKKINDA BİLGİ")
        b3 = sds.get("b3_bilesim") or {}
        tip = b3.get("tip") or "karisim"

        if tip == "madde":
            madde = b3.get("madde") or {}
            cls._add_subsection_title(doc, s3_t.get("s3_1") or "3.1. Maddeler")
            cls._add_kv_line(doc, s3_t.get("chemical_identity") or "Kimyasal Kimliği", madde.get("kimyasal_kimlik") or "—")
            cls._add_kv_line(doc, "CAS & EC No", f"CAS: {madde.get('cas_no') or '—'} | EC: {madde.get('ec_no') or '—'}")
        else:
            karisim = b3.get("karisim") or {}
            bilesenler = karisim.get("bilesenler") or []
            cls._add_subsection_title(doc, s3_t.get("s3_2") or "3.2. Karışımlar")
            if bilesenler:
                tbl = doc.add_table(rows=len(bilesenler)+1, cols=6)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl_t = s3_t.get("table", {})
                headers = [
                    tbl_t.get("comp_name") or "Bileşen Adı",
                    tbl_t.get("cas_no") or "CAS No",
                    tbl_t.get("ec_no") or "EC No",
                    tbl_t.get("reg_no") or "Kayıt No",
                    tbl_t.get("concentration") or "Konsantrasyon",
                    tbl_t.get("classification") or "Sınıflandırma"
                ]
                for i, h in enumerate(headers):
                    cell = tbl.cell(0, i)
                    set_cell_borders(cell, color="444444", sz="4")
                    set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                    p = cell.paragraphs[0]
                    r = p.add_run(h)
                    r.font.name = FONT_NAME
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.size = Pt(10)

                for r_idx, b in enumerate(bilesenler, start=1):
                    row = tbl.rows[r_idx]
                    vals = [
                        b.get("ad", ""),
                        b.get("cas_no", ""),
                        b.get("ec_no", ""),
                        b.get("kayit_no", ""),
                        b.get("konsantrasyon", ""),
                        b.get("siniflandirma", "")
                    ]
                    for c_idx, val in enumerate(vals):
                        c = row.cells[c_idx]
                        set_cell_borders(c, color="444444", sz="4")
                        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
                        p = c.paragraphs[0]
                        r_v = p.add_run(str(val or "—"))
                        r_v.font.name = FONT_NAME
                        r_v.font.size = Pt(8)

                doc.add_paragraph()
            else:
                cls._add_text_line(doc, s3_t.get("no_hazardous_ingredients") or "Zararlı sınır değerini aşan bileşen bulunmamaktadır.")

    @classmethod
    def _add_section_4(cls, doc, sds, t, lang):
        s4_t = t.get("s4", {})
        cls._add_section_banner(doc, s4_t.get("title") or "4. İLK YARDIM ÖNLEMLERİ")
        b4 = sds.get("b4_ilk_yardim") or {}
        b4_1 = b4.get("b4_1") or {}

        cls._add_subsection_title(doc, s4_t.get("s4_1") or "4.1. İlk Yardım Önlemlerinin Açıklanması")
        cls._add_kv_line(doc, s4_t.get("inhalation") or "Solunması Halinde", b4_1.get("soluma") or ("Remove casualty to fresh air." if lang == "en" else "Kazazedeyi temiz havaya çıkarın."))
        cls._add_kv_line(doc, s4_t.get("skin") or "Cilt ile Teması Halinde", b4_1.get("cilt_temasi") or ("Wash with plenty of soap and water." if lang == "en" else "Bol su ve sabun ile yıkayınız."))
        cls._add_kv_line(doc, s4_t.get("eye") or "Göz ile Teması Halinde", b4_1.get("goz_temasi") or ("Rinse thoroughly with plenty of water for at least 15 minutes." if lang == "en" else "Bol su ile en az 15 dakika yıkayınız."))
        cls._add_kv_line(doc, s4_t.get("ingestion") or "Yutulması Halinde", b4_1.get("yutma") or ("Rinse mouth with water. Do NOT induce vomiting." if lang == "en" else "Ağzı su ile çalkalayınız. Kusturmayınız."))

        cls._add_subsection_title(doc, s4_t.get("s4_2") or "4.2. Akut ve Sonradan Görülen En Önemli Belirtiler ve Etkiler")
        cls._add_text_line(doc, b4.get("b4_2_belirtiler_etkiler") or ("No significant symptoms known." if lang == "en" else "Önemli bir belirti bildirilmemiştir."))

        cls._add_subsection_title(doc, s4_t.get("s4_3") or "4.3. Tıbbi Müdahale ve Özel Tedavi Gereği İçin İlk İşaretler")
        cls._add_text_line(doc, b4.get("b4_3_acil_tibbi_mudahale") or (s4_t.get("symptomatic_treatment") or "Semptomatik tedavi uygulayınız."))

    @classmethod
    def _add_section_5(cls, doc, sds, t, lang):
        s5_t = t.get("s5", {})
        cls._add_section_banner(doc, s5_t.get("title") or "5. YANGINLA MÜCADELE ÖNLEMLERİ")
        b5 = sds.get("b5_yangin_mucadele") or {}
        b5_1 = b5.get("b5_1") or {}

        cls._add_subsection_title(doc, s5_t.get("s5_1") or "5.1. Yangın Söndürücüler")
        cls._add_kv_line(doc, s5_t.get("suitable_media") or "Uygun Yangın Söndürücüler", b5_1.get("uygun_sondurucu") or ("Foam, dry chemical powder, CO2, water spray." if lang == "en" else "Köpük, kuru kimyevi toz, CO2, su sisi."))
        cls._add_kv_line(doc, s5_t.get("unsuitable_media") or "Uygun Olmayan Söndürücüler", b5_1.get("uygun_olmayan_sondurucu") or ("High pressure full water jet." if lang == "en" else "Yüksek basınçlı tam su jeti."))

        cls._add_subsection_title(doc, s5_t.get("s5_2") or "5.2. Madde veya Karışımdan Kaynaklanan Özel Zararlar")
        cls._add_text_line(doc, b5.get("b5_2_ozel_zararlar") or ("Thermal decomposition may release toxic carbon oxides." if lang == "en" else "Yanma halinde toksik karbon oksitler açığa çıkabilir."))

        cls._add_subsection_title(doc, s5_t.get("s5_3") or "5.3. Yangın Söndürme Ekipleri İçin Tavsiyeler")
        cls._add_text_line(doc, b5.get("b5_3_sondurme_ekibi_tavsiyeleri") or (s5_t.get("fire_protection") or "Tam koruyucu teçhizat ve solunum cihazı kullanınız."))

    @classmethod
    def _add_section_6(cls, doc, sds, t, lang):
        s6_t = t.get("s6", {})
        cls._add_section_banner(doc, s6_t.get("title") or "6. KAZA SONUCU YAYILMAYA KARŞI ÖNLEMLER")
        b6 = sds.get("b6_kaza_sonucu_yayilma") or {}
        b6_1 = b6.get("b6_1") or {}

        cls._add_subsection_title(doc, s6_t.get("s6_1") or "6.1. Kişisel Önlemler, Koruyucu Donanım ve Acil Durum Prosedürleri")
        cls._add_text_line(doc, b6_1.get("kisisel_onlemler_acil_olmayan") or ("Ventilate area. Keep away from ignition sources." if lang == "en" else "Alanı havalandırın. Ateş kaynaklarını uzaklaştırın."))

        cls._add_subsection_title(doc, s6_t.get("s6_2") or "6.2. Çevresel Önlemler")
        cls._add_text_line(doc, b6.get("b6_2_cevresel_onlemler") or ("Prevent entry into drains and waterways." if lang == "en" else "Kanalizasyon ve su yollarına karışmasını önleyiniz."))

        cls._add_subsection_title(doc, s6_t.get("s6_3") or "6.3. Muhafaza Etme ve Temizleme İçin Yöntemler ve Materyaller")
        cls._add_text_line(doc, b6.get("b6_3_kontrol_temizleme_yontemleri") or ("Absorb with inert material (sand, silica gel)." if lang == "en" else "İnert emici materyal (kum vb.) ile toplayınız."))

        cls._add_subsection_title(doc, s6_t.get("s6_4") or "6.4. Diğer Bölümlere Atıflar")
        cls._add_text_line(doc, b6.get("b6_4_diger_bolumlere_atif") or (s6_t.get("see_sections") or "Kişisel korunma için Bölüm 8'e, bertaraf için Bölüm 13'e bakınız."))

    @classmethod
    def _add_section_7(cls, doc, sds, t, lang):
        s7_t = t.get("s7", {})
        cls._add_section_banner(doc, s7_t.get("title") or "7. ELLEÇLEME VE DEPOLAMA")
        b7 = sds.get("b7_ellecme_depolama") or {}
        b7_2 = b7.get("b7_2") or {}

        cls._add_subsection_title(doc, s7_t.get("s7_1") or "7.1. Güvenli Elleçleme İçin Önlemler")
        cls._add_text_line(doc, b7.get("b7_1_guvenli_ellecleme") or ("Provide adequate ventilation. Avoid contact." if lang == "en" else "İyi havalandırılan yerlerde kullanın. Temastan kaçının."))

        cls._add_subsection_title(doc, s7_t.get("s7_2") or "7.2. Uyuşmazlıkları da İçeren Güvenli Depolama Koşulları")
        cls._add_text_line(doc, b7_2.get("guvenli_depolama_kosullari") or ("Store in a cool, dry, well-ventilated place." if lang == "en" else "Serin, kuru, iyi havalandırılan yerde saklayınız."))

        cls._add_subsection_title(doc, s7_t.get("s7_3") or "7.3. Belirli Son Kullanımlar")
        cls._add_text_line(doc, b7.get("b7_3_belirli_son_kullanimlar") or (s7_t.get("no_extra_uses") or "Bölüm 1.2'de belirtilen alanlar içindir."))

    @classmethod
    def _add_section_8(cls, doc, sds, t, lang):
        s8_t = t.get("s8", {})
        cls._add_section_banner(doc, s8_t.get("title") or "8. MARUZ KALMA KONTROLLERİ / KİŞİSEL KORUNMA")
        b8 = sds.get("b8_maruz_kalma_kontrolu") or {}
        b8_1 = b8.get("b8_1_kontrol_parametreleri") or []
        b8_2 = b8.get("b8_2") or {}
        kkd = b8_2.get("kkd") or {}

        cls._add_subsection_title(doc, s8_t.get("s8_1") or "8.1. Kontrol Parametreleri")
        if b8_1:
            tbl = doc.add_table(rows=len(b8_1)+1, cols=3)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_t = s8_t.get("table", {})
            headers = [
                tbl_t.get("substance") or "Madde / Bileşen",
                tbl_t.get("limit_value") or "Mesleki Maruziyet Sınır Değeri",
                tbl_t.get("legal_basis") or "Yasal Dayanak"
            ]
            for i, h in enumerate(headers):
                cell = tbl.cell(0, i)
                set_cell_borders(cell, color="444444", sz="4")
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                p = cell.paragraphs[0]
                r = p.add_run(h)
                r.font.name = FONT_NAME
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.size = Pt(10)

            for r_idx, exp in enumerate(b8_1, start=1):
                row = tbl.rows[r_idx]
                vals = [exp.get("madde", ""), exp.get("sinir_degeri", ""), exp.get("yasal_dayanak", "")]
                for c_idx, val in enumerate(vals):
                    c = row.cells[c_idx]
                    set_cell_borders(c, color="444444", sz="4")
                    set_cell_margins(c, top=60, bottom=60, left=80, right=80)
                    p = c.paragraphs[0]
                    r_v = p.add_run(str(val or "—"))
                    r_v.font.name = FONT_NAME
                    r_v.font.size = Pt(8)

            doc.add_paragraph()
        else:
            cls._add_text_line(doc, s8_t.get("no_limits") or "Tanımlı mesleki maruziyet sınır değeri bulunmamaktadır.")

        cls._add_subsection_title(doc, s8_t.get("s8_2") or "8.2. Maruz Kalma Kontrolleri")
        cls._add_kv_line(doc, s8_t.get("engineering_controls") or "Uygun Mühendislik Kontrolleri", b8_2.get("muhendislik_kontrolleri") or ("Provide adequate ventilation." if lang == "en" else "Yeterli genel ve lokal havalandırma sağlayınız."))
        cls._add_kv_line(doc, s8_t.get("eye_protection") or "Göz / Yüz Koruması", kkd.get("goz_yuz") or ("Safety glasses with side-shields (EN 166)." if lang == "en" else "EN 166 uyumlu koruyucu gözlük."))
        cls._add_kv_line(doc, s8_t.get("hand_protection") or "Ellerin Korunması", kkd.get("cilt_el") or ("Chemical resistant gloves (EN 374)." if lang == "en" else "EN 374 uyumlu nitril eldiven."))
        cls._add_kv_line(doc, s8_t.get("respiratory_protection") or "Solunum Sisteminin Korunması", kkd.get("solunum") or ("Use suitable respirator with filter type A in case of insufficient ventilation." if lang == "en" else "Gerekli hallerde A tipi filtreli solunum maskesi."))

    @classmethod
    def _add_section_9(cls, doc, sds, t, lang):
        s9_t = t.get("s9", {})
        cls._add_section_banner(doc, s9_t.get("title") or "9. FİZİKSEL VE KİMYASAL ÖZELLİKLER")
        b9 = sds.get("b9_fiziksel_kimyasal_ozellikler") or {}
        b9_1 = b9.get("b9_1") or {}
        p_t = s9_t.get("props", {})

        cls._add_subsection_title(doc, s9_t.get("s9_1") or "9.1. Temel Fiziksel ve Kimyasal Özellikler Hakkında Bilgi")
        props = [
            (p_t.get("appearance") or "a) Görünüm (Fiziksel hal, renk)", b9_1.get("gorunum")),
            (p_t.get("odour") or "b) Koku", b9_1.get("koku") or ("Characteristic" if lang == "en" else "Karakteristik")),
            (p_t.get("odour_threshold") or "c) Koku Eşiği", b9_1.get("koku_esigi")),
            (p_t.get("ph") or "ç) pH", b9_1.get("ph")),
            (p_t.get("melting_point") or "d) Erime / Donma Noktası", b9_1.get("erime_noktasi")),
            (p_t.get("boiling_point") or "e) İlk Kaynama Noktası ve Kaynama Aralığı", b9_1.get("kaynama_noktasi")),
            (p_t.get("flash_point") or "f) Parlama Noktası", b9_1.get("parlama_noktasi")),
            (p_t.get("evaporation_rate") or "g) Buharlaşma Hızı", b9_1.get("buharlasma_hizi")),
            (p_t.get("flammability") or "ğ) Alevlenirlik (Katı, Gaz)", b9_1.get("alevlenirlik")),
            (p_t.get("explosive_limits") or "h) Üst / Alt Alevlenirlik veya Patlayıcı Limitleri", b9_1.get("ust_alt_limitler")),
            (p_t.get("vapour_pressure") or "ı) Buhar Basıncı", b9_1.get("buhar_basinci")),
            (p_t.get("vapour_density") or "i) Buhar Yoğunluğu", b9_1.get("buhar_yogunlugu")),
            (p_t.get("relative_density") or "j) Bağıl Yoğunluk", b9_1.get("bagil_yogunluk")),
            (p_t.get("solubility") or "k) Çözünürlük", b9_1.get("cozunurluk")),
            (p_t.get("partition_coeff") or "l) Dağılım Katsayısı (n-oktanol/su)", b9_1.get("dagilim_katsayisi_log_kow")),
            (p_t.get("auto_ignition") or "m) Kendiliğinden Tutuşma Sıcaklığı", b9_1.get("kendiliginden_tutusma_sicakligi")),
            (p_t.get("decomposition") or "n) Bozunma Sıcaklığı", b9_1.get("bozunma_sicakligi")),
            (p_t.get("viscosity") or "o) Akışkanlık (Viskozite)", b9_1.get("akiskanlik")),
            (p_t.get("explosive_props") or "ö) Patlayıcı Özellikler", b9_1.get("patlayici_ozellikler") or (s9_t.get("not_explosive") or "Patlayıcı değildir.")),
            (p_t.get("oxidising_props") or "p) Oksitleyici Özellikler", b9_1.get("oksitleyici_ozellikler") or (s9_t.get("not_oxidising") or "Oksitleyici değildir."))
        ]
        default_empty = s9_t.get("not_determined") or "Belirlenmemiştir"
        for k, v in props:
            cls._add_kv_line(doc, k, v or default_empty)

        cls._add_subsection_title(doc, s9_t.get("s9_2") or "9.2. Diğer Bilgiler")
        cls._add_text_line(doc, b9.get("b9_2_diger_bilgiler") or ("No additional information." if lang == "en" else "Ek bilgi bulunmamaktadır."))

    @classmethod
    def _add_section_10(cls, doc, sds, t, lang):
        s10_t = t.get("s10", {})
        cls._add_section_banner(doc, s10_t.get("title") or "10. KARARLILIK VE TEPKİME")
        b10 = sds.get("b10_kararlilik_tepkime") or {}

        cls._add_subsection_title(doc, s10_t.get("s10_1") or "10.1. Tepkime")
        cls._add_text_line(doc, b10.get("b10_1_tepkime") or (s10_t.get("no_reactivity") or "Normal koşullarda tehlikeli bir reaksiyon vermez."))

        cls._add_subsection_title(doc, s10_t.get("s10_2") or "10.2. Kimyasal Kararlılık")
        cls._add_text_line(doc, b10.get("b10_2_kimyasal_kararlilik") or (s10_t.get("stable") or "Normal depolama ve kullanım koşullarında kararlıdır."))

        cls._add_subsection_title(doc, s10_t.get("s10_3") or "10.3. Zararlı Tepkime Olasılığı")
        cls._add_text_line(doc, b10.get("b10_3_zararli_reaksiyon_olasiligi") or (s10_t.get("no_hazardous_reactions") or "Tehlikeli reaksiyon beklenmez."))

        cls._add_subsection_title(doc, s10_t.get("s10_4") or "10.4. Kaçınılması Gereken Durumlar")
        cls._add_text_line(doc, b10.get("b10_4_kacinilmasi_gereken_durumlar") or ("Excessive heat, sparks, open flame." if lang == "en" else "Aşırı ısı, kıvılcım, açık alev."))

        cls._add_subsection_title(doc, s10_t.get("s10_5") or "10.5. Kaçınılması Gereken Maddeler")
        cls._add_text_line(doc, b10.get("b10_5_kacinilmasi_gereken_maddeler") or ("Strong acids, bases and oxidising agents." if lang == "en" else "Kuvvetli asitler, bazlar ve oksitleyiciler."))

        cls._add_subsection_title(doc, s10_t.get("s10_6") or "10.6. Zararlı Bozunma Ürünleri")
        cls._add_text_line(doc, b10.get("b10_6_zararli_bozunma_urunleri") or (s10_t.get("no_decomp") or "Normal depolamada ayrışmaz. Yangında toksik gazlar açığa çıkar."))

    @classmethod
    def _add_section_11(cls, doc, sds, t, lang):
        s11_t = t.get("s11", {})
        cls._add_section_banner(doc, s11_t.get("title") or "11. TOKSİKOLOJİK BİLGİLER")
        b11 = sds.get("b11_toksikolojik") or {}
        b11_1 = b11.get("b11_1") or {}

        cls._add_subsection_title(doc, s11_t.get("s11_1") or "11.1. Toksik Etkiler Hakkında Bilgi")
        cls._add_kv_line(doc, s11_t.get("acute_tox") or "a) Akut Toksisite", b11_1.get("akut_toksisite") or ("Does not meet classification criteria." if lang == "en" else "Kriterleri karşılamamaktadır."))
        cls._add_kv_line(doc, s11_t.get("skin_corr") or "b) Cilt Aşınması / Tahrişi", b11_1.get("cilt_asinmasi_tahrisi") or ("Not classified based on available data." if lang == "en" else "Mevcut bilgilere göre sınıflandırılmaz."))
        cls._add_kv_line(doc, s11_t.get("eye_damage") or "c) Ciddi Göz Hasarları / Tahrişi", b11_1.get("goz_hasari") or ("Not classified based on available data." if lang == "en" else "Mevcut bilgilere göre sınıflandırılmaz."))
        cls._add_kv_line(doc, s11_t.get("sensitisation") or "ç) Solunum Yolları veya Cilt Hassaslaşması", b11_1.get("solunum_cilt_hassasiyeti") or ("No sensitising effect known." if lang == "en" else "Hassaslaştırıcı etkisi bildirilmemiştir."))
        cls._add_kv_line(doc, s11_t.get("mutagenicity") or "d) Eşey Hücre Mutajenitesi", b11_1.get("mutajenite") or ("Not classified as mutagenic." if lang == "en" else "Mutajenik olarak sınıflandırılmaz."))
        cls._add_kv_line(doc, s11_t.get("carcinogenicity") or "e) Kanserojenite", b11_1.get("kanserojenite") or ("Contains no carcinogenic substances." if lang == "en" else "Kanserojen madde içermez."))
        cls._add_kv_line(doc, s11_t.get("repr_tox") or "f) Üreme Toksisitesi", b11_1.get("ureme_toksisitesi") or ("Not toxic for reproduction." if lang == "en" else "Üreme için toksik değildir."))
        cls._add_kv_line(doc, s11_t.get("stot_se") or "g) Belirli Hedef Organ Toksisitesi (BHOT) - Tek Maruz Kalma", b11_1.get("bhot_tek_maruz") or ("Does not meet criteria." if lang == "en" else "Kriterleri karşılamaz."))
        cls._add_kv_line(doc, s11_t.get("stot_re") or "ğ) Belirli Hedef Organ Toksisitesi (BHOT) - Tekrarlı Maruz Kalma", b11_1.get("bhot_tekrarli") or ("No target organ damage expected." if lang == "en" else "Organ hasarı beklenmez."))
        cls._add_kv_line(doc, s11_t.get("aspiration") or "h) Aspirasyon Zararı", b11_1.get("aspirasyon_zarari") or ("No aspiration hazard expected." if lang == "en" else "Aspirasyon tehlikesi oluşturmaz."))

    @classmethod
    def _add_section_12(cls, doc, sds, t, lang):
        s12_t = t.get("s12", {})
        cls._add_section_banner(doc, s12_t.get("title") or "12. EKOLOJİK BİLGİLER")
        b12 = sds.get("b12_ekolojik") or {}

        cls._add_subsection_title(doc, s12_t.get("s12_1") or "12.1. Toksisite")
        cls._add_text_line(doc, b12.get("b12_1_toksisite") or ("Not classified as environmentally hazardous." if lang == "en" else "Sucul organizmalar için zararlı olarak sınıflandırılmamıştır."))

        cls._add_subsection_title(doc, s12_t.get("s12_2") or "12.2. Kalıcılık ve Bozunabilirlik")
        cls._add_text_line(doc, b12.get("b12_2_kalicilik_bozunabilirlik") or ("Components are biodegradable." if lang == "en" else "Bileşenleri biyolojik olarak ayrışabilir."))

        cls._add_subsection_title(doc, s12_t.get("s12_3") or "12.3. Biyobirikim Potansiyeli")
        cls._add_text_line(doc, b12.get("b12_3_biyobirikim") or ("Low bioaccumulation potential." if lang == "en" else "Biyobirikim potansiyeli düşüktür."))

        cls._add_subsection_title(doc, s12_t.get("s12_4") or "12.4. Toprakta Hareketlilik")
        cls._add_text_line(doc, b12.get("b12_4_topraktaki_hareketlilik") or ("Mobility depends on water solubility." if lang == "en" else "Topraktaki hareketliliği çözünürlüğe bağlıdır."))

        cls._add_subsection_title(doc, s12_t.get("s12_5") or "12.5. PBT ve vPvB Değerlendirmesinin Sonuçları")
        cls._add_text_line(doc, b12.get("b12_5_pbt_vpvb_sonuclari") or (t.get("s2", {}).get("pbt_vpvb") or "PBT / vPvB kriterlerini karşılamaz."))

        cls._add_subsection_title(doc, s12_t.get("s12_7") or "12.6. Diğer Olumsuz Etkiler")
        cls._add_text_line(doc, b12.get("b12_6_diger_olumsuz_etkiler") or ("No other adverse environmental effects known." if lang == "en" else "Bilinen başka bir olumsuz etkisi yoktur."))

    @classmethod
    def _add_section_13(cls, doc, sds, t, lang):
        s13_t = t.get("s13", {})
        cls._add_section_banner(doc, s13_t.get("title") or "13. BERTARAF ETME BİLGİLERİ")
        b13 = sds.get("b13_bertaraf") or {}

        cls._add_subsection_title(doc, s13_t.get("s13_1") or "13.1. Atık İşleme Yöntemleri")
        cls._add_kv_line(doc, "Atık İşleme", b13.get("b13_1_atik_isleme_yontemleri") or (s13_t.get("waste_advice") or "Ulusal mevzuata uygun olarak lisanslı atık bertaraf tesislerine verilmelidir."))
        cls._add_kv_line(doc, "Ambalaj Atıkları", b13.get("b13_1_ambalaj_atik_isleme") or ("Empty packaging must be handed over to licensed recovery facilities." if lang == "en" else "Boş ambalajlar lisanslı geri kazanım/bertaraf firmalarına teslim edilmelidir."))
        cls._add_kv_line(doc, "Kanalizasyon Uyarısı", b13.get("b13_1_kanalizasyon_uyarisi") or ("Do not discharge into drains or waterways." if lang == "en" else "Kanalizasyona ve su kaynaklarına dökülmemelidir."))

    @classmethod
    def _add_section_14(cls, doc, sds, t, lang):
        s14_t = t.get("s14", {})
        cls._add_section_banner(doc, s14_t.get("title") or "14. TAŞIMACILIK BİLGİSİ")
        b14 = sds.get("b14_tasimacilik") or {}

        cls._add_subsection_title(doc, s14_t.get("s14_1") or "14.1. UN Numarası")
        cls._add_text_line(doc, b14.get("b14_1_un_numarasi") or (s14_t.get("not_dangerous_goods") or "Taşımacılıkta tehlikeli madde değildir."))

        cls._add_subsection_title(doc, s14_t.get("s14_2") or "14.2. Uygun UN Taşımacılık Adı")
        cls._add_text_line(doc, b14.get("b14_2_un_tasimacilik_adi") or ("Not applicable" if lang == "en" else "Uygulanabilir değildir."))

        cls._add_subsection_title(doc, s14_t.get("s14_3") or "14.3. Taşımacılık Zararlılık Sınıf(lar)ı")
        cls._add_text_line(doc, b14.get("b14_3_tasimacilik_sinifi") or "—")

        cls._add_subsection_title(doc, s14_t.get("s14_4") or "14.4. Ambalajlama Grubu")
        cls._add_text_line(doc, b14.get("b14_4_ambalajlama_grubu") or "—")

        cls._add_subsection_title(doc, s14_t.get("s14_5") or "14.5. Çevresel Zararlar")
        cls._add_text_line(doc, b14.get("b14_5_cevresel_zararlar") or ("Not a Marine Pollutant" if lang == "en" else "Deniz Kirletici değildir."))

        cls._add_subsection_title(doc, s14_t.get("s14_6") or "14.6. Kullanıcı İçin Özel Önlemler")
        cls._add_text_line(doc, b14.get("b14_6_kullanici_ozel_onlemler") or ("Secure containers against falling and damage during transport." if lang == "en" else "Taşımada devrilme ve hasara karşı emniyete alınız."))

        cls._add_subsection_title(doc, s14_t.get("s14_7") or "14.7. MARPOL 73/78 Ek II ve IBC Koduna Göre Dökme Taşımacılık")
        cls._add_text_line(doc, b14.get("b14_7_marpol_ibc") or ("Not applicable (Packaged goods)." if lang == "en" else "Uygulanabilir değildir (Ambalajlı sevkiyat)."))

    @classmethod
    def _add_section_15(cls, doc, sds, t, lang):
        s15_t = t.get("s15", {})
        cls._add_section_banner(doc, s15_t.get("title") or "15. MEVZUAT BİLGİSİ")
        b15 = sds.get("b15_mevzuat") or {}

        cls._add_subsection_title(doc, s15_t.get("s15_1") or "15.1. Madde veya Karışıma Özel Güvenlik, Sağlık ve Çevre Mevzuatı")
        cls._add_text_line(doc, b15.get("b15_1_ozel_mevzuat_hukumleri") or ("Regulation (EC) No 1907/2006 (REACH), Regulation (EC) No 1272/2008 (CLP)." if lang == "en" else "KKDİK Yönetmeliği (RG: 30105), SEA Yönetmeliği (RG: 28848)."))

        cls._add_subsection_title(doc, s15_t.get("s15_2") or "15.2. Kimyasal Güvenlik Değerlendirmesi")
        cls._add_text_line(doc, b15.get("b15_2_kimyasal_guvenlik_degerlendirmesi") or (s15_t.get("csa_not_carried_out") or "Bu karışım için Kimyasal Güvenlik Değerlendirmesi (KGD) yapılmamıştır."))

    @classmethod
    def _add_section_16(cls, doc, sds, t, lang):
        s16_t = t.get("s16", {})
        cls._add_section_banner(doc, s16_t.get("title") or "16. DİĞER BİLGİLER")
        b16 = sds.get("b16_diger_bilgiler") or {}
        
        if lang == "en":
            found_h_codes = reference_service.extract_h_codes_from_sds(sds)
            tam_h = translation_service.format_h_statements(found_h_codes, lang="en")
        else:
            tam_h = b16.get("tam_h_ifadeleri") or []

        cls._add_subsection_title(doc, s16_t.get("h_statements_full") or "16.1. Zararlılık (H) İfadelerinin Tam Metinleri")
        if tam_h:
            for h in tam_h:
                cls._add_text_line(doc, f"• {h}")
        else:
            cls._add_text_line(doc, "H-ifadesi bulunmamaktadır.")

        cls._add_subsection_title(doc, "16.5. Eğitim Tavsiyeleri")
        cls._add_text_line(doc, b16.get("egitim_tavsiyeleri") or "Çalışanlar kimyasalların güvenli elleçlenmesi ve KKD kullanımı konusunda eğitilmelidir.")
