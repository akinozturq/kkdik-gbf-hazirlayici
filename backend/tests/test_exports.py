import io
import pytest
from docx import Document
from fastapi.testclient import TestClient

def test_export_docx_endpoint(client: TestClient):
    # 1. Create a test product
    create_res = client.post("/api/products", json={
        "urun_adi": "Polyester Reçine PRS-100",
        "ticari_kod": "AYP-PRS-100",
        "kategori": "Reçineler",
        "sds_data": {
            "meta": {"hazirlama_tarihi": "18.08.2026", "revizyon_no": "01"},
            "b1_kimlik": {
                "b1_1": {"madde_karisim_adi": "Polyester Reçine PRS-100"},
                "b1_3": {"tedarikci_adi": "Aypol Kimya"},
                "b1_4": {"acil_telefon": "114"}
            },
            "b2_zarar_tanimi": {
                "b2_1": {"siniflandirmalar": [{"zararlilik_sinifi": "Alevlenir Sıvı", "h_kodu": "H226"}]},
                "b2_2": {"piktogramlar": ["GHS02"], "uyari_kelimesi": "Dikkat", "h_ifadeleri": ["H226"]}
            }
        }
    })
    assert create_res.status_code == 201
    prod_id = create_res.json()["id"]

    # 2. Test DOCX export
    docx_res = client.get(f"/api/products/{prod_id}/export/docx")
    assert docx_res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in docx_res.headers["content-type"]
    assert len(docx_res.content) > 1000

    doc = Document(io.BytesIO(docx_res.content))
    # Header check
    header_tables = doc.sections[0].header.tables
    assert len(header_tables) == 1
    h_row = header_tables[0].rows[0]
    assert len(h_row.cells) == 3
    assert "GÜVENLİK BİLGİ FORMU" in h_row.cells[1].text
    assert "Bu belge, 23 Haziran 2017 tarihli ve 30105 sayılı Resmi Gazete’de yayımlanan" in h_row.cells[1].text
    assert "Hazırlama Tarihi" in h_row.cells[2].text and "18.08.2026" in h_row.cells[2].text
    assert "Revizyon No" in h_row.cells[2].text and "01" in h_row.cells[2].text

    # Footer check
    footer_tables = doc.sections[0].footer.tables
    assert len(footer_tables) == 1
    row = footer_tables[0].rows[0]
    assert len(row.cells) == 4
    assert "POLCHEM" in row.cells[0].text and "www.polchem.com.tr" in row.cells[0].text
    assert "AYPOL" in row.cells[1].text and "www.aypol.com.tr" in row.cells[1].text
    assert "GÖKAY" in row.cells[2].text and "www.gokayboya.com.tr" in row.cells[2].text
    assert "Sayfa" in row.cells[3].text

def test_export_pdf_endpoint(client: TestClient):
    # 1. Create a test product
    create_res = client.post("/api/products", json={
        "urun_adi": "Selülozik Tiner Extra",
        "ticari_kod": "G.TN.20.00",
        "kategori": "Solventler & Tinerler",
        "sds_data": {
            "meta": {"hazirlama_tarihi": "18.08.2026"},
            "b1_kimlik": {
                "b1_1": {"madde_karisim_adi": "Selülozik Tiner Extra"},
                "b1_3": {"tedarikci_adi": "Aypol Kimya"},
                "b1_4": {"acil_telefon": "114"}
            }
        }
    })
    assert create_res.status_code == 201
    prod_id = create_res.json()["id"]

    # 2. Test PDF export
    pdf_res = client.get(f"/api/products/{prod_id}/export/pdf")
    assert pdf_res.status_code == 200
    assert "application/pdf" in pdf_res.headers["content-type"]
    assert pdf_res.content.startswith(b"%PDF")
    assert len(pdf_res.content) > 1000

def test_preview_html_endpoint(client: TestClient):
    create_res = client.post("/api/products", json={
        "urun_adi": "Epoksi Astar",
        "ticari_kod": "AYP-EPX-01",
        "kategori": "Epoksi Sistemler",
        "sds_data": {}
    })
    prod_id = create_res.json()["id"]

    html_res = client.get(f"/api/products/{prod_id}/export/preview-html")
    assert html_res.status_code == 200
    assert "text/html" in html_res.headers["content-type"]
    assert "GÜVENLİK BİLGİ FORMU" in html_res.text
    assert "Bu belge, 23 Haziran 2017 tarihli ve 30105 sayılı Resmi Gazete’de yayımlanan" in html_res.text
    assert "Epoksi Astar" in html_res.text

    # Test English HTML Preview (REACH Annex II)
    html_en_res = client.get(f"/api/products/{prod_id}/export/preview-html?lang=en")
    assert html_en_res.status_code == 200
    assert "SAFETY DATA SHEET" in html_en_res.text
    assert "Identification of the substance/mixture" in html_en_res.text
    assert "Hazards identification" in html_en_res.text


def test_export_english_docx_and_pdf(client: TestClient):
    create_res = client.post("/api/products", json={
        "urun_adi": "Endüstriyel Tiner",
        "ticari_kod": "SOL-200",
        "kategori": "Solventler",
        "sds_data": {
            "meta": {"hazirlama_tarihi": "20.08.2026", "revizyon_no": "02"},
            "b1_kimlik": {
                "b1_1": {"madde_karisim_adi": "Endüstriyel Tiner", "kayit_numarasi": "Kayıttan muaftır / Uygulanabilir değildir."},
                "b1_2": {"tanimlanmis_kullanimlar": ["Sanayi / Endüstriyel kullanım"]},
                "b1_3": {"tedarikci_adi": "Aypol Kimya", "adres": "İstanbul / Türkiye"},
                "b1_4": {"acil_telefon": "114 (UZEM - Ulusal Zehir Danışma Merkezi)"}
            },
            "b2_zarar_tanimi": {
                "b2_1": {"siniflandirmalar": [{"zararlilik_sinifi": "Alevlenir Sıvı", "kategori": "Kategori 2", "h_kodu": "H225"}]},
                "b2_2": {"piktogramlar": ["GHS02"], "uyari_kelimesi": "Tehlike", "h_ifadeleri": ["H225"]},
                "b2_3": {"pbt_vpvb_degerlendirme": "PBT / vPvB kriterlerini karşılamaz."}
            },
            "b4_ilk_yardim": {
                "b4_1": {"soluma": "Kazazedeyi temiz havaya çıkarın.", "cilt_temasi": "Bol su ve sabun ile yıkayınız."}
            }
        }
    })
    prod_id = create_res.json()["id"]

    # English HTML Preview Content Translation Verification
    html_en_res = client.get(f"/api/products/{prod_id}/export/preview-html?lang=en")
    assert html_en_res.status_code == 200
    assert "SAFETY DATA SHEET" in html_en_res.text
    assert "Flammable Liquid" in html_en_res.text
    assert "Danger" in html_en_res.text
    assert "Remove casualty to fresh air" in html_en_res.text
    assert "Wash thoroughly with plenty of soap and water" in html_en_res.text
    assert "Industrial / Professional use" in html_en_res.text
    assert "Istanbul / Turkey" in html_en_res.text

    # English DOCX
    docx_en_res = client.get(f"/api/products/{prod_id}/export/docx?lang=en")
    assert docx_en_res.status_code == 200
    assert len(docx_en_res.content) > 1000

    doc_en = Document(io.BytesIO(docx_en_res.content))
    header_en_tables = doc_en.sections[0].header.tables
    assert len(header_en_tables) == 1
    assert "SAFETY DATA SHEET" in header_en_tables[0].rows[0].cells[1].text
    assert "According to Regulation (EC) No. 1907/2006" in header_en_tables[0].rows[0].cells[1].text
    assert "Compilation Date" in header_en_tables[0].rows[0].cells[2].text

    footer_en_tables = doc_en.sections[0].footer.tables
    assert len(footer_en_tables) == 1
    row_en = footer_en_tables[0].rows[0]
    assert "POLCHEM" in row_en.cells[0].text
    assert "Page" in row_en.cells[3].text

    # English PDF
    pdf_en_res = client.get(f"/api/products/{prod_id}/export/pdf?lang=en")
    assert pdf_en_res.status_code == 200
    assert pdf_en_res.content.startswith(b"%PDF")
    assert len(pdf_en_res.content) > 1000


def test_export_404_for_missing_product(client: TestClient):
    res = client.get("/api/products/99999/export/docx")
    assert res.status_code == 404
